"""Tables: extraction from PDFs / spreadsheets, a text form for retrieval, and
deterministic answering of numeric table questions.

Why this module exists
----------------------
A RAG pipeline that sends a table to a language model and asks it to "add up the
March invoices" will sometimes get it wrong, and it will never be able to show its
working. Tables are the one place where the right answer is computable, so here it
is computed - in Python, with integers and floats, never by the model.

Three jobs, in order:

1. *Extraction* - `extract_pdf_tables` (PyMuPDF's ruling/whitespace table finder)
   and `extract_sheet_tables` (.xlsx via openpyxl, .csv/.tsv via the csv module)
   both produce the same plain dict shape that `Store.replace_tables` stores:
   ``{"page", "title", "headers", "rows", "numeric_cols"}``.

2. *Retrievability* - `table_to_text` / `table_to_chunks` render a table as compact
   markdown-ish text including its title and header row, so a table is still found
   by ordinary vector and FTS retrieval ("what did we pay ABC Traders") even when no
   arithmetic is involved.

3. *Computation* - `answer_table_query` reads a question, works out the operation
   (sum / average / count / max / min / difference / plain filter), the numeric
   column it applies to and the filters that narrow the rows, then computes the
   value and returns it together with the rows it used and a one-line `workings`
   string. When it cannot work any of that out with confidence it returns ``None``
   and the caller falls back to normal retrieval - a wrong number is much worse
   than no number.

Everything here is a pure function. Tables are passed in by the caller (usually
from `Store.list_tables`), extra keys such as ``id`` / ``doc_id`` / ``doc_name``
ride along untouched and are copied into the answer so the UI can cite the source.
No network, no model, no module-level mutable state.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Iterable, Sequence

# ------------------------------------------------------------------ constants
MAX_SHEET_ROWS = 5000          # per sheet: a runaway spreadsheet must not blow up the index
MAX_CHUNK_CHARS = 1500         # one table chunk stays comfortably inside the embedder window
MAX_CELL_CHARS = 80            # a single essay-in-a-cell must not eat a whole chunk
MAX_MATCHED_ROWS = 50          # evidence rows returned to the caller
NUMERIC_COL_RATIO = 0.60       # >= this share of non-empty cells must parse as numbers
SHEET_EXT = {".xlsx", ".xlsm", ".xltx"}
DELIMITED_EXT = {".csv", ".tsv"}

# Cells that mean "no value" rather than a number. Compared lower-cased.
NULLISH = frozenset({
    "", "-", "--", "---", "–", "—", ".", "..", "n/a", "na", "n.a", "n.a.",
    "nil", "none", "null", "nan", "tbd", "tba", "?", "--", "not applicable",
})

_CURRENCY_SYMBOLS = "₹$€£¥"                     # rupee dollar euro pound yen
_CURRENCY_WORDS = re.compile(r"\b(?:inr|rs|rupees?|usd|eur|gbp|aed|sgd)\b\.?", re.I)
# A number core: optional sign, then digits with ',' grouping (Indian 1,50,000 or 150,000),
# then an optional decimal part.
_NUM_CORE = re.compile(r"^([+-]?)(\d[\d,]*)(?:\.(\d+))?")
# Anything allowed to trail a number and still leave it a number: units and percent.
_UNIT_TAIL = re.compile(r"^[A-Za-zµ°%/·^²³.\-\s]{1,14}$")
# Indian magnitude words are part of the number, not a unit.
_MULTIPLIERS = {"lakh": 1e5, "lakhs": 1e5, "lac": 1e5, "lacs": 1e5,
                "crore": 1e7, "crores": 1e7, "cr": 1e7}

_MONTH_NAMES = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "oct": 10, "october": 10,
    "nov": 11, "november": 11, "dec": 12, "december": 12,
}
_MONTH_LABELS = ["", "January", "February", "March", "April", "May", "June",
                 "July", "August", "September", "October", "November", "December"]

# Words that carry the question's intent rather than naming data. Kept out of table
# scoring and out of text-filter candidates so "total spend" cannot filter on a
# vendor literally called "Total".
_INTENT_WORDS = frozenset("""total sum summed combined altogether aggregate overall average mean avg count
number numbers how much many highest largest biggest maximum max lowest smallest minimum min least difference
change delta variance compare comparison list show which what who when where give me all any the of in for on
at by from to with and or is are was were do did does we our us their my me please tell value values amount amounts
rows row table tables column columns per across between above below over under more less than most spend spent
paid pay cost costs price prices""".split())

_STOP = _INTENT_WORDS | frozenset("""a an as be been being it its this that these those there here about into
if then so such not no nor but only also very can could should would shall will may might must have has had
each other others same both few new old year years month months date dates day days
""".split())

# Question words that point at a money column, and at a quantity column.
_MONEY_SYNONYMS = frozenset("""spend spent spending expenditure expense expenses amount amt cost costs price
prices value paid pay pays paying payment payments total revenue sales salary fee fees charge charges billing
billed bill bills invoiced purchase purchases purchased money rupees inr rs budget turnover""".split())
_QTY_SYNONYMS = frozenset("""qty quantity quantities units unit nos no pieces pcs volume count number
headcount strength""".split())
# Header words that make a column money-ish / quantity-ish.
_MONEY_HEADER_WORDS = frozenset("""amount amt cost price value total spend expenditure expense invoice revenue
sales salary fee fees charge charges payment paid billing gst tax rs inr rupees inr. budget premium""".split())
_QTY_HEADER_WORDS = frozenset("""qty quantity units unit nos pieces pcs count number strength headcount""".split())
_MONEY_MARKER = re.compile(r"[₹$€£]|\b(?:inr|rs\.?|rupees?)\b", re.I)
# Question-side markers that force money formatting (per the spec: rupee, INR, Rs, amount).
_MONEY_QUESTION = re.compile(r"[₹]|\b(?:inr|rs\.?|rupees?|amount)\b", re.I)

# Header words that mark a column as carrying a date or a period.
_DATE_HEADER_WORDS = frozenset("""date dates month months period periods quarter quarters year years fy
billing posted posting when day week invoice_date duedate due""".split())

# Operations, checked in this order. "how many" must beat "total", "average" must
# beat "total", so the narrower intents come first.
_OP_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("max", re.compile(r"\b(?:max(?:imum)?|highest|largest|biggest|greatest|peak|most\s+expensive|dearest|top|"
                       r"tallest|longest|heaviest|oldest|latest|fastest)\b", re.I)),
    ("min", re.compile(r"\b(?:min(?:imum)?|lowest|smallest|cheapest|least\s+expensive|least|shortest|"
                       r"lightest|earliest|slowest)\b", re.I)),
    # "mean" only counts as an average when it reads like a noun: "what does E-47
    # mean?" is not a request for an average.
    ("avg", re.compile(r"\b(?:average|avg)\b|\bmean\s+(?:of|value)\b|\bthe\s+mean\b|"
                       r"\bper\s+(?:month|invoice|row|item|unit|vendor)\b", re.I)),
    ("diff", re.compile(r"\b(?:difference|differences|delta|variance|spread|range|gap|change)\b", re.I)),
    ("count", re.compile(r"\b(?:how\s+many|count|number\s+of|no\.?\s+of|nos\.?\s+of)\b", re.I)),
    ("sum", re.compile(r"\b(?:total|totals|sum|summed|combined|altogether|aggregate|overall|"
                       r"how\s+much|spend|spent|spending|expenditure|paid|billed)\b", re.I)),
)
# A bare lookup ("list invoices from ABC") when no aggregation word is present.
_LOOKUP_INTENT = re.compile(r"\b(?:list|show|display|find|give|which|what|who|all|every|details?\s+of)\b", re.I)

_BETWEEN_RE = re.compile(
    r"\bbetween\s+(?P<lo>[₹$€£]?\s*\d[\d,]*(?:\.\d+)?(?:\s*(?:lakh|lakhs|lac|lacs|crore|crores|%))?)"
    r"\s*(?:and|to|-|–)\s*"
    r"(?P<hi>[₹$€£]?\s*\d[\d,]*(?:\.\d+)?(?:\s*(?:lakh|lakhs|lac|lacs|crore|crores|%))?)", re.I)
_NUM_TOKEN = r"(?P<num>[₹$€£]?\s*\d[\d,]*(?:\.\d+)?(?:\s*(?:lakh|lakhs|lac|lacs|crore|crores|%))?)"
_COMPARE_RES: tuple[tuple[str, re.Pattern[str]], ...] = (
    (">=", re.compile(r"\b(?:at\s+least|no\s+less\s+than|minimum\s+of|from)\s+" + _NUM_TOKEN, re.I)),
    ("<=", re.compile(r"\b(?:at\s+most|no\s+more\s+than|up\s+to|maximum\s+of)\s+" + _NUM_TOKEN, re.I)),
    (">", re.compile(r"\b(?:above|more\s+than|greater\s+than|over|exceed(?:s|ing)?|higher\s+than|bigger\s+than|"
                     r">=?)\s*" + _NUM_TOKEN, re.I)),
    ("<", re.compile(r"\b(?:below|less\s+than|under|fewer\s+than|lower\s+than|smaller\s+than|cheaper\s+than|"
                     r"<=?)\s*" + _NUM_TOKEN, re.I)),
)

_MONTH_QUESTION_RE = re.compile(
    r"\b(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|"
    r"sep(?:t|tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\b", re.I)
_YEAR_QUESTION_RE = re.compile(r"\b((?:19|20)\d{2})\b")
_FY_QUESTION_RE = re.compile(r"\bFY\s*[-'’]?\s*(\d{4}|\d{2})\b", re.I)
_QUARTER_QUESTION_RE = re.compile(r"\bQ([1-4])\b", re.I)

# Cheap router check: does this question look like it wants a number out of a table?
_TABLE_QUESTION_RES: tuple[re.Pattern[str], ...] = (
    re.compile(r"\b(?:total|totals|sum|subtotal|aggregate|altogether|combined|overall)\b", re.I),
    re.compile(r"\b(?:average|avg|median)\b|\bmean\s+(?:of|value)\b|\bthe\s+mean\b", re.I),
    re.compile(r"\bhow\s+(?:much|many)\b", re.I),
    re.compile(r"\b(?:count|number\s+of|no\.?\s+of)\b", re.I),
    re.compile(r"\b(?:highest|largest|biggest|max(?:imum)?|lowest|smallest|min(?:imum)?|cheapest|top)\b", re.I),
    re.compile(r"\b(?:spend|spent|spending|expenditure|invoiced|billed|paid)\b", re.I),
    re.compile(r"\b(?:compare|difference|breakdown|per\s+(?:month|vendor|item))\b", re.I),
    re.compile(r"\b(?:above|below|more\s+than|less\s+than|greater\s+than|between)\s+"
               r"[₹$]?\s*\d", re.I),
    re.compile(r"\b(?:in|for|during)\s+(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
               r"jul(?:y)?|aug(?:ust)?|sep(?:t|tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?|"
               r"Q[1-4]|FY\s*\d{2,4})\b", re.I),
)


# ============================================================ number parsing
def parse_number(cell: str) -> float | None:
    """Parse one table cell into a float, or ``None`` when it is not a number.

    Indian and Western money and measurement notation both appear in the same
    documents, so this is deliberately forgiving in the ways that are unambiguous
    and strict everywhere else:

    * ``"1,50,000"`` / ``"150,000"`` -> 150000.0   (commas are grouping separators)
    * ``"₹1,20,000.50"``, ``"Rs. 4,500"``, ``"INR 900"`` -> the number, symbol dropped
    * ``"(1,200)"`` -> -1200.0                     (accounting parentheses mean negative)
    * ``"12%"`` -> 12.0                            (the percent sign is a unit, not /100)
    * ``"40 Nm"``, ``"132 mg/dL"`` -> 40.0, 132.0  (a trailing unit is dropped)
    * ``"2 lakh"``, ``"1.5 crore"`` -> 200000.0, 15000000.0
    * ``"-"``, ``"N/A"``, ``"nil"``, ``""`` -> None (an empty cell, not a zero)
    * ``"1,200 2,300"``, ``"approx 500"`` -> None  (two numbers, or unknown prefix)

    Returning ``None`` rather than 0.0 for a blank matters: a blank must not drag an
    average down, and must not count towards the numeric-column ratio.
    """
    if cell is None:
        return None
    s = str(cell).replace(" ", " ").replace(" ", " ").strip()
    if s.lower() in NULLISH:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    s = _CURRENCY_WORDS.sub(" ", s)
    s = "".join(ch for ch in s if ch not in _CURRENCY_SYMBOLS).strip()
    if s.startswith("."):
        s = "0" + s
    match = _NUM_CORE.match(s)
    if not match:
        return None
    sign, whole, frac = match.group(1), match.group(2), match.group(3)
    digits = whole.replace(",", "")
    if not digits.isdigit():
        return None
    try:
        value = float(digits + ("." + frac if frac else ""))
    except ValueError:
        return None
    tail = s[match.end():].strip()
    if tail:
        key = tail.lower().strip(" .")
        if key in _MULTIPLIERS:
            value *= _MULTIPLIERS[key]
        elif not _UNIT_TAIL.match(tail):
            return None           # something that is not a unit follows: not a number
    if sign == "-" or negative:
        value = -value
    return value


_EMPTY_CELLS = {"-", "--", "---", "\u2013", "\u2014", "n/a", "na", "nil", "none", "\u2014\u2014", "."}


_INDEX_HEADERS = frozenset({"#", "no", "no.", "sno", "s.no", "s no", "sl", "sl.", "sr", "sr.",
                            "serial", "index", "item no", "row", "rank"})


_UNIT_SUFFIX = re.compile(r"^[^A-Za-z]*[\d.,]+\s*([A-Za-z%°/µ·]{0,8})\.?$")


def _is_quantity(cell: str) -> bool:
    """Is this cell a measured value rather than a date or a label?

    "18 cm", "9.25%", "₹57,530" and "36" are quantities. "03 Sep" and "March 2026" are dates:
    the letters that follow the number name a month, not a unit.
    """
    text = _squash(cell)
    if not text or parse_number(text) is None:
        return False
    m = _UNIT_SUFFIX.match(text)
    if not m:
        return False
    unit = m.group(1).lower().strip(".")
    return unit not in _MONTH_NAMES


def _is_index_column(name: str, values: Sequence[str]) -> bool:
    """A column of 1, 2, 3 … is a row number: summing it means nothing."""
    if _norm_text(name).strip() in _INDEX_HEADERS:
        return True
    numbers = [parse_number(v) for v in values]
    numbers = [n for n in numbers if n is not None]
    return len(numbers) >= 3 and numbers == [float(i + 1) for i in range(len(numbers))]


def detect_numeric_cols(headers: Sequence[str], rows: Sequence[Sequence[str]],
                        min_ratio: float = NUMERIC_COL_RATIO) -> list[str]:
    """Names of the columns that hold numbers.

    A column counts as numeric when at least `min_ratio` (default 60%) of its
    *non-empty* cells parse with `parse_number`. The threshold is below 100% on
    purpose: real tables carry "-", "N/A" and the odd footnote in an otherwise
    numeric column, and those cells should not disqualify it. Columns that are
    entirely empty are never numeric.
    """
    numeric: list[str] = []
    for index, name in enumerate(headers):
        filled = 0
        parsed = 0
        for row in rows:
            cell = _squash(row[index] if index < len(row) else "")
            # "-", "—", "N/A" and friends mean "nothing in this cell", so they must not
            # count against the column. A Credit column that is mostly dashes with three
            # real amounts in it is still a numeric column.
            if not cell or cell.lower() in _EMPTY_CELLS:
                continue
            filled += 1
            if _is_quantity(cell):
                parsed += 1
        if filled and parsed and parsed / filled >= min_ratio:
            values = [_squash(row[index] if index < len(row) else "") for row in rows]
            if not _is_index_column(name, values):
                numeric.append(name)
    return numeric


# ============================================================ grid utilities
def _squash(cell: Any) -> str:
    """One table cell -> single-line, collapsed-whitespace text."""
    if cell is None:
        return ""
    return re.sub(r"\s+", " ", str(cell).replace("­", "")).strip()


def _clean_grid(grid: Iterable[Sequence[Any]]) -> list[list[str]]:
    """Clean every cell, then drop rows and columns that are entirely empty.

    PDF table finders routinely emit a blank spacer column or a blank rule row;
    keeping them would add phantom columns named "col3" and wreck column matching.
    """
    cleaned = [[_squash(c) for c in row] for row in grid]
    if not cleaned:
        return []
    width = max((len(r) for r in cleaned), default=0)
    if not width:
        return []
    cleaned = [r + [""] * (width - len(r)) for r in cleaned]
    keep_cols = [i for i in range(width) if any(r[i] for r in cleaned)]
    if not keep_cols:
        return []
    return [[r[i] for i in keep_cols] for r in cleaned if any(r[i] for i in keep_cols)]


def _looks_like_header(first: Sequence[str], rest: Sequence[Sequence[str]]) -> bool:
    """Is this first row a header row, or is it data?

    Header rows are labels: mostly filled, none of them a number, none of them a
    sentence. A first row that contains a number is data, and its table gets
    synthesised "col1..colN" headers instead - guessing a header from data would
    silently mislabel every answer built on it.
    """
    filled = [c for c in first if c]
    if not filled:
        return False
    if any(parse_number(c) is not None for c in filled):
        return False
    if len(filled) < max(1, (len(first) + 1) // 2):
        return False
    if all(len(c) > 60 for c in filled):
        return False
    return True


def _name_headers(header_row: Sequence[str]) -> list[str]:
    """Fill in blank header cells and make duplicates distinct.

    Answers are addressed by column *name*, so names have to exist and be unique.
    """
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, raw in enumerate(header_row):
        name = _squash(raw)[:MAX_CELL_CHARS] or f"col{index + 1}"
        key = name.lower()
        if key in seen:
            seen[key] += 1
            name = f"{name}_{seen[key]}"
        else:
            seen[key] = 1
        names.append(name)
    return names


def _finalise_table(page: int, title: str, grid: Sequence[Sequence[str]],
                    force_header: bool = False) -> dict | None:
    """Turn a cleaned grid into the table dict the store and the rest of this module use."""
    grid = _clean_grid(grid)
    if not grid:
        return None
    if (force_header or _looks_like_header(grid[0], grid[1:])) and len(grid) > 1:
        headers = _name_headers(grid[0])
        rows = [list(r) for r in grid[1:]]
    else:
        headers = [f"col{i + 1}" for i in range(len(grid[0]))]
        rows = [list(r) for r in grid]
    return {
        "page": int(page),
        "title": _squash(title),
        "headers": headers,
        "rows": rows,
        "numeric_cols": detect_numeric_cols(headers, rows),
    }


# ============================================================ PDF extraction
def _heading_ish(line: str) -> bool:
    """A short, label-like line - the sort of thing that captions a table.

    A caption is a title, not the tail of a paragraph, so a line that ends in a full stop,
    starts mid-sentence, or reads like prose ("... and some should trigger abstention.")
    is rejected. A table with no caption is better titled "Table on page 4" than with a
    sentence fragment.
    """
    line = line.strip()
    if not (3 <= len(line) <= 70):
        return False
    if line.endswith((".", ",", ";")):          # a sentence, not a caption
        return False
    if line.endswith(":"):
        line = line[:-1].strip()
    if not re.search(r"[A-Za-z]", line):
        return False
    words = line.split()
    if len(words) > 8:
        return False
    if words and words[0][:1].islower():        # continues the sentence above it
        return False
    return parse_number(line) is None


def _table_title(page: Any, bbox: Sequence[float]) -> str:
    """Nearest heading-ish line above the table on the same page, else "".

    Cheap and best effort, exactly as specified: one `get_text("blocks")` call, take
    the closest block that ends above the table and whose last line reads like a
    caption. A table with no caption gets "" rather than a wrong one.
    """
    try:
        top = float(bbox[1])
        best_text = ""
        best_bottom = -1.0
        for block in page.get_text("blocks"):
            y1, text = float(block[3]), str(block[4])
            if y1 > top + 1 or top - y1 > 120:      # below the table, or too far above
                continue
            for line in reversed(text.splitlines()):
                line = _squash(line)
                if _heading_ish(line):
                    if y1 > best_bottom:
                        best_bottom, best_text = y1, line
                    break
        return best_text
    except Exception:
        return ""


def extract_pdf_tables(path: Path) -> list[dict]:
    """Every table PyMuPDF can find in a PDF, page by page.

    `page.find_tables()` locates tables from ruling lines, falling back to column
    whitespace. Its own header detection is used when it reports an *external*
    header (a label row drawn above the table body, which `extract()` leaves out);
    otherwise the first extracted row is judged by `_looks_like_header`.

    A page whose table finder raises is skipped rather than failing the document -
    the same best-effort stance `ingest.py` takes with OCR.
    """
    import pymupdf as fitz

    out: list[dict] = []
    with fitz.open(path) as doc:
        if doc.needs_pass:
            raise ValueError("PDF is password protected")
        for index, page in enumerate(doc):
            try:
                found = page.find_tables()
            except Exception:
                continue
            for table in getattr(found, "tables", []) or []:
                try:
                    raw = [list(r) for r in table.extract()]
                except Exception:
                    continue
                prefix: list[list[Any]] = []
                header = getattr(table, "header", None)
                names = list(getattr(header, "names", None) or []) if header is not None else []
                if header is not None and getattr(header, "external", False) and any(_squash(n) for n in names):
                    prefix = [names]                     # a header row drawn outside the body
                built = _finalise_table(
                    page=index + 1,
                    title=_table_title(page, getattr(table, "bbox", (0, 0, 0, 0))),
                    grid=prefix + raw,
                    force_header=bool(prefix),
                )
                if built and built["rows"]:
                    out.append(built)
    return out


# ==================================================== spreadsheet extraction
def _cell_to_str(value: Any) -> str:
    """One spreadsheet value -> text, keeping it recognisable to `parse_number`.

    Floats that are whole numbers lose their ".0" (openpyxl reads every integer as a
    float), and dates become ISO text so the period filters can read them back.
    """
    if value is None:
        return ""
    if value is True or value is False:
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value == int(value) and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if hasattr(value, "strftime"):
        if getattr(value, "hour", 0) or getattr(value, "minute", 0):
            return value.strftime("%Y-%m-%d %H:%M")
        return value.strftime("%Y-%m-%d")
    return _squash(value)


def _sniff_delimiter(sample: str, suffix: str) -> str:
    """Delimiter for a delimited text file: sniff it, fall back on the extension."""
    default = "\t" if suffix == ".tsv" else ","
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except Exception:
        return default


def extract_sheet_tables(path: Path) -> list[dict]:
    """Tables from a spreadsheet: one table per sheet.

    * ``.xlsx`` / ``.xlsm`` - openpyxl in ``read_only`` + ``data_only`` mode, so a
      large workbook is streamed rather than loaded, and formula cells give their
      last cached *value* instead of "=SUM(B2:B9)".
      ``page`` is the sheet's position (1-based) and ``title`` is the sheet name,
      which keeps citations pointing at something a person can find.
    * ``.csv`` / ``.tsv`` - the stdlib csv reader with a sniffed delimiter,
      ``page`` 1 and ``title`` the file stem.

    Each sheet is capped at `MAX_SHEET_ROWS` data rows: past a few thousand rows a
    table stops being something to quote and becomes a database, and the cap keeps
    indexing bounded.
    """
    suffix = path.suffix.lower()
    if suffix in DELIMITED_EXT:
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
        delimiter = _sniff_delimiter(raw[:4096], suffix)
        grid = [[_squash(c) for c in row]
                for row in csv.reader(raw.splitlines(), delimiter=delimiter)][: MAX_SHEET_ROWS + 1]
        built = _finalise_table(page=1, title=path.stem, grid=grid)
        return [built] if built and built["rows"] else []

    if suffix not in SHEET_EXT:
        raise ValueError(f"Not a spreadsheet: {path.suffix}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:      # pragma: no cover - depends on the install
        raise RuntimeError("openpyxl is required to read .xlsx files") from exc

    out: list[dict] = []
    book = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        for index, sheet in enumerate(book.worksheets):
            grid: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                grid.append([_cell_to_str(c) for c in row])
                if len(grid) > MAX_SHEET_ROWS:       # header row + MAX_SHEET_ROWS of data
                    break
            built = _finalise_table(page=index + 1, title=str(sheet.title), grid=grid)
            if built and built["rows"]:
                out.append(built)
    finally:
        book.close()
    return out


def extract_tables(path: Path) -> list[dict]:
    """Dispatch on file type; an unsupported type has no tables rather than raising."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_tables(path)
    if suffix in SHEET_EXT or suffix in DELIMITED_EXT:
        return extract_sheet_tables(path)
    return []


# ============================================================ text rendering
def _short(cell: str, limit: int = MAX_CELL_CHARS) -> str:
    cell = _squash(cell)
    return cell if len(cell) <= limit else cell[: limit - 1].rstrip() + "…"


def _row_line(cells: Sequence[str], width: int) -> str:
    padded = [_short(c) for c in cells] + [""] * max(0, width - len(cells))
    return "| " + " | ".join(padded[:width]) + " |"


def _table_head(t: dict) -> str:
    """Title line + markdown header + separator: repeated at the top of every chunk."""
    headers = [str(h) for h in t.get("headers") or []]
    page = t.get("page")
    title = _squash(t.get("title") or "")
    label = f"Table: {title}" if title else "Table"
    if t.get("doc_name"):
        label += f" ({t['doc_name']}"
        label += f", page {page})" if page else ")"
    elif page:
        label += f" (page {page})"
    width = len(headers)
    return "\n".join([label, "", _row_line(headers, width), "| " + " | ".join(["---"] * width) + " |"])


def _packed_chunks(t: dict, max_chars: int) -> list[tuple[int, str]]:
    """Pack a table's rows into chunks: [(rows in this chunk, chunk text), ...].

    Shared by the two public renderers so they always agree on how many rows the
    first chunk actually shows.
    """
    headers = [str(h) for h in t.get("headers") or []]
    rows = [list(r) for r in t.get("rows") or []]
    head = _table_head(t)
    if not rows:
        return [(0, head)]
    width = len(headers)
    total = len(rows)
    # Both renderers append a short note - "(rows 31-60 of 400)" here, "... 340 more
    # rows" in `table_to_text` - so room for them is reserved up front and no chunk
    # can come out over budget.
    reserve = len(f"\n(rows {total}-{total} of {total})\n… {total} more rows")
    budget = max(120, max_chars - reserve)
    out: list[tuple[int, str]] = []
    index = 0
    while index < total:
        first = index
        body: list[str] = []
        size = len(head) + 1
        while index < total:
            line = _row_line(rows[index], width)
            if body and size + len(line) + 1 > budget:
                break
            body.append(line)
            size += len(line) + 1
            index += 1
        note = "" if (first == 0 and index == total) else f"\n(rows {first + 1}-{index} of {total})"
        out.append((index - first, head + note + "\n" + "\n".join(body)))
    return out


def table_to_chunks(t: dict, max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    """The whole table as retrieval text, split into chunks of at most `max_chars`.

    Every chunk repeats the title line and the header row, so a chunk lifted out of
    the middle of a long table is still self-describing - both to the embedder and
    to a person reading the citation. Continuation chunks say which rows they cover.
    A single row longer than the whole budget still gets its own chunk: dropping
    data silently would be worse than one oversized chunk.
    """
    return [text for _, text in _packed_chunks(t, max_chars)]


def table_to_text(t: dict, max_chars: int = MAX_CHUNK_CHARS) -> str:
    """One compact markdown-ish rendering of a table for embedding and keyword search.

    Title, header row and as many data rows as fit in `max_chars`; if the table is
    longer, the text ends with a "... N more rows" note and `table_to_chunks` is the
    way to get all of it. Keeping the title and headers in the text is the whole
    point - it is what lets ordinary retrieval find a table at all.
    """
    packed = _packed_chunks(t, max_chars)
    shown, text = packed[0]
    rest = max(0, len(t.get("rows") or []) - shown)
    return text + (f"\n… {rest} more rows" if rest else "")


# ============================================================ question intent
def is_table_question(question: str) -> bool:
    """Cheap router check: does this question want a number out of a table?

    Deliberately generous, because it only decides whether to *try*
    `answer_table_query`, which then refuses on its own when it cannot compute an
    answer. False positives cost one pass of regexes; false negatives cost a wrong
    answer from a language model.
    """
    q = str(question or "")
    return any(pattern.search(q) for pattern in _TABLE_QUESTION_RES)


def _words(text: str) -> list[str]:
    """Lower-cased alphanumeric words, e.g. "Amount (INR)" -> ["amount", "inr"]."""
    return [w for w in re.split(r"[^0-9a-z₹]+", str(text or "").lower()) if w]


def _norm_text(text: str) -> str:
    """Lower-cased text with punctuation flattened to single spaces, for word-boundary tests."""
    return " " + " ".join(_words(text)) + " "


def _parse_operation(question: str) -> str | None:
    """The arithmetic the question asks for, or "filter" for a plain lookup, or None."""
    for name, pattern in _OP_PATTERNS:
        if pattern.search(question):
            return name
    if _LOOKUP_INTENT.search(question):
        return "filter"
    return None


# ============================================================ column matching
def _col_index(headers: Sequence[str], name: str) -> int:
    for index, header in enumerate(headers):
        if str(header) == str(name):
            return index
    return -1


def _cell(row: Sequence[str], index: int) -> str:
    return str(row[index]) if 0 <= index < len(row) else ""


def _is_money_header(header: str) -> bool:
    words = set(_words(header))
    if words & _QTY_HEADER_WORDS:
        return False                        # "Units"/"Qty" is a count, never money
    return bool(words & _MONEY_HEADER_WORDS) or bool(_MONEY_MARKER.search(str(header)))


def _column_is_money(headers: Sequence[str], rows: Sequence[Sequence[str]], name: str,
                     question: str) -> bool:
    """Should this column's values be printed as rupees?

    True when the header names money, when the cells themselves carry a currency
    marker, or when the question asks in money terms (rupee sign, INR, Rs, "amount").
    A quantity header vetoes all of that.
    """
    index = _col_index(headers, name)
    if index < 0:
        return False
    if set(_words(name)) & _QTY_HEADER_WORDS:
        return False
    if _is_money_header(name):
        return True
    marked = 0
    filled = 0
    for row in rows:
        cell = _cell(row, index)
        if not cell.strip():
            continue
        filled += 1
        if _MONEY_MARKER.search(cell):
            marked += 1
    if filled and marked / filled >= 0.3:
        return True
    return bool(_MONEY_QUESTION.search(question))


def _resolve_numeric_column(question: str, headers: Sequence[str],
                            rows: Sequence[Sequence[str]],
                            numeric_cols: Sequence[str]) -> tuple[str | None, int]:
    """Pick the numeric column the question is about. Returns (name, score).

    Scoring, highest first:
      +10  the whole header appears in the question ("total invoice amount")
      +4   per significant header word found in the question
      +3   the question uses a money/quantity synonym and the header is that kind

    A tie between two different columns resolves to ``None``: "average unit price"
    over a table with both *Price* and *Units* is genuinely ambiguous, and guessing
    is exactly what this module must not do. When nothing matches at all but the
    table has a single numeric column, that column is used with a low score, which
    the caller reports as ``"medium"`` confidence - there is no other candidate to
    confuse it with.
    """
    qwords = set(_words(question))
    qnorm = _norm_text(question)
    money_asked = bool(qwords & _MONEY_SYNONYMS)
    qty_asked = bool(qwords & _QTY_SYNONYMS)

    scored: list[tuple[int, str]] = []
    for name in numeric_cols:
        words = _words(name)
        if not words:
            continue
        score = 0
        phrase = " ".join(words)
        if len(phrase) >= 3 and f" {phrase} " in qnorm:
            score += 10
        def matches(word: str) -> bool:
            """Exact word, or the same stem: "credited"/"credits" both mean Credit."""
            if word in qwords:
                return True
            stem = word[:max(4, len(word) - 2)]
            return any(q == word or q.startswith(stem) for q in qwords if len(q) >= 4)

        score += 4 * sum(1 for w in words if len(w) >= 3 and w not in _STOP and matches(w))
        header_words = set(words)
        if money_asked and (header_words & _MONEY_HEADER_WORDS or _MONEY_MARKER.search(str(name))):
            score += 3
        if qty_asked and header_words & _QTY_HEADER_WORDS:
            score += 3
        if score:
            scored.append((score, name))

    if scored:
        scored.sort(key=lambda pair: (-pair[0], pair[1]))
        if len(scored) > 1 and scored[0][0] == scored[1][0]:
            return None, 0                      # ambiguous: refuse rather than guess
        return scored[0][1], scored[0][0]

    if len(numeric_cols) == 1:
        return numeric_cols[0], 1
    return None, 0


def _money_numeric_column(headers: Sequence[str], rows: Sequence[Sequence[str]],
                          numeric_cols: Sequence[str]) -> str | None:
    """The one money-ish numeric column, if there is exactly one.

    Used for a bare comparison such as "invoices above 50000", where the question
    names no column: over a table of *Vendor / Qty / Amount* only *Amount* can be
    meant. With zero or several money columns this returns ``None`` and the
    comparison is dropped, which in turn makes the whole answer refuse.
    """
    money = [c for c in numeric_cols if _is_money_header(c)]
    return money[0] if len(money) == 1 else None


# ================================================================== periods
def _four_digit_year(value: int | str) -> int:
    year = int(value)
    if year < 100:
        year += 2000 if year < 70 else 1900
    return year


def _year_month_of(cell: str) -> tuple[int | None, int | None]:
    """Best-effort (year, month) for a date-ish cell; either part may be None.

    Handles ISO ``2026-03-15``, ``15/03/2026`` and ``03/2026``, month names with or
    without a year (``March``, ``Mar-26``, ``March 2026``) and a bare year.
    Two-digit day/month pairs are read **day first** (``15/03/2026``), the Indian
    convention, unless the first number cannot be a day.
    """
    text = _squash(cell)
    if not text:
        return None, None

    iso = re.search(r"\b((?:19|20)\d{2})[-/.](\d{1,2})(?:[-/.](\d{1,2}))?\b", text)
    if iso:
        month = int(iso.group(2))
        return _four_digit_year(iso.group(1)), month if 1 <= month <= 12 else None

    dmy = re.search(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})\b", text)
    if dmy:
        first, second = int(dmy.group(1)), int(dmy.group(2))
        month = second if first > 12 or second <= 12 else first
        return _four_digit_year(dmy.group(3)), month if 1 <= month <= 12 else None

    named = _MONTH_QUESTION_RE.search(text)
    if named:
        month = _MONTH_NAMES.get(named.group(1).lower())
        year_match = re.search(r"[-/'’\s,]\s*((?:19|20)?\d{2})\b", text[named.end():])
        year = _four_digit_year(year_match.group(1)) if year_match else None
        return year, month

    my = re.search(r"\b(\d{1,2})[-/](\d{4})\b", text)
    if my:
        month = int(my.group(1))
        return _four_digit_year(my.group(2)), month if 1 <= month <= 12 else None

    bare_year = re.fullmatch(r"\s*((?:19|20)\d{2})(?:\.0)?\s*", text)
    if bare_year:
        return _four_digit_year(bare_year.group(1)), None
    return None, None


def _period_bounds(op: str, value: str) -> tuple[tuple[int, int] | None, tuple[int, int] | None,
                                                 frozenset[int]]:
    """A period filter value -> inclusive (year, month) bounds plus the months it allows.

    Both are returned because cells are inconsistent: ``2026-03-15`` can be checked
    against the bounds, while a bare ``March`` can only be checked against the month
    set. Fiscal periods follow the Indian year, April to March, so ``FY26`` is
    April 2025 - March 2026 and ``Q1 FY26`` is April-June 2025. A plain ``Q1`` with
    no ``FY`` is the calendar quarter, January-March.
    """
    text = str(value or "")
    if op == "month":
        month = 0
        found = _MONTH_QUESTION_RE.search(text)
        if found:
            month = _MONTH_NAMES.get(found.group(1).lower(), 0)
        year_match = _YEAR_QUESTION_RE.search(text)
        if month and year_match:
            year = _four_digit_year(year_match.group(1))
            return (year, month), (year, month), frozenset({month})
        return None, None, frozenset({month} if month else ())
    if op == "year":
        year_match = _YEAR_QUESTION_RE.search(text)
        if not year_match:
            return None, None, frozenset()
        year = _four_digit_year(year_match.group(1))
        return (year, 1), (year, 12), frozenset(range(1, 13))
    if op == "fy":
        fy = _FY_QUESTION_RE.search(text)
        if not fy:
            return None, None, frozenset()
        end = _four_digit_year(fy.group(1))
        return (end - 1, 4), (end, 3), frozenset(range(1, 13))
    if op == "quarter":
        quarter = _QUARTER_QUESTION_RE.search(text)
        if not quarter:
            return None, None, frozenset()
        index = int(quarter.group(1))
        fy = _FY_QUESTION_RE.search(text)
        if fy:
            end = _four_digit_year(fy.group(1))
            start_month = 4 + 3 * (index - 1)
            if start_month <= 12:
                start = (end - 1, start_month)
                finish = (end - 1, start_month + 2) if start_month + 2 <= 12 else (end, start_month + 2 - 12)
            else:
                start = (end, start_month - 12)
                finish = (end, start_month - 12 + 2)
            months = frozenset(((start_month - 1 + k) % 12) + 1 for k in range(3))
            return start, finish, months
        year_match = _YEAR_QUESTION_RE.search(text)
        start_month = 1 + 3 * (index - 1)
        months = frozenset(range(start_month, start_month + 3))
        if year_match:
            year = _four_digit_year(year_match.group(1))
            return (year, start_month), (year, start_month + 2), months
        return None, None, months
    return None, None, frozenset()


def _match_period(cell: str, op: str, value: str) -> bool:
    """Does a date-ish cell fall inside a period filter?

    When the filter names a year and the cell carries one, both year and month must
    agree. When the cell has no year - a column of bare month names - only the month
    is checked, which is the best that data supports.
    """
    low, high, months = _period_bounds(op, value)
    year, month = _year_month_of(cell)
    if op == "year":
        return year is not None and low is not None and year == low[0]
    if month is None:
        return False
    if year is not None and low is not None and high is not None:
        key = year * 12 + month
        return low[0] * 12 + low[1] <= key <= high[0] * 12 + high[1]
    return month in months


def _period_column(headers: Sequence[str], rows: Sequence[Sequence[str]],
                   exclude: Sequence[str] = ()) -> str | None:
    """The column that carries dates or periods, chosen by header wording and content."""
    best: tuple[float, str] | None = None
    for index, name in enumerate(headers):
        if name in exclude:
            continue
        header_hit = bool(set(_words(name)) & _DATE_HEADER_WORDS)
        filled = 0
        dated = 0
        for row in rows:
            cell = _cell(row, index)
            if not cell.strip():
                continue
            filled += 1
            year, month = _year_month_of(cell)
            if year is not None or month is not None:
                dated += 1
        share = dated / filled if filled else 0.0
        if not header_hit and share < 0.5:
            continue
        score = share + (1.0 if header_hit else 0.0)
        if best is None or score > best[0]:
            best = (score, name)
    return best[1] if best else None


def _period_filter(question: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
                   exclude: Sequence[str] = ()) -> dict | None:
    """Read a month / year / quarter / fiscal-year filter out of the question."""
    month = _MONTH_QUESTION_RE.search(question)
    year = _YEAR_QUESTION_RE.search(question)
    fy = _FY_QUESTION_RE.search(question)
    quarter = _QUARTER_QUESTION_RE.search(question)
    if not (month or year or fy or quarter):
        return None
    column = _period_column(headers, rows, exclude=exclude)
    if column is None:
        return None
    if quarter:
        label = f"Q{quarter.group(1)}"
        if fy:
            label += f" FY{fy.group(1)}"
        elif year:
            label += f" {year.group(1)}"
        return {"column": column, "op": "quarter", "value": label}
    if month:
        name = _MONTH_LABELS[_MONTH_NAMES[month.group(1).lower()]]
        return {"column": column, "op": "month",
                "value": f"{name} {year.group(1)}" if year else name}
    if fy:
        return {"column": column, "op": "fy", "value": f"FY{fy.group(1)}"}
    return {"column": column, "op": "year", "value": year.group(1)}


# ================================================================== filters
def _numeric_filters(question: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
                     numeric_cols: Sequence[str], target: str | None) -> tuple[list[dict], bool]:
    """Comparisons such as "above 50000" or "between 10 and 20".

    Returns (filters, ok). ``ok`` is False when a comparison was clearly asked for
    but no column could be attached to it - the caller then refuses the whole
    question instead of quietly answering a different one.
    """
    filters: list[dict] = []
    wanted: list[tuple[str, Any]] = []

    between = _BETWEEN_RE.search(question)
    if between:
        low = parse_number(between.group("lo"))
        high = parse_number(between.group("hi"))
        if low is not None and high is not None:
            wanted.append(("between", [min(low, high), max(low, high)]))
    else:
        for op, pattern in _COMPARE_RES:
            found = pattern.search(question)
            if not found:
                continue
            value = parse_number(found.group("num"))
            if value is not None:
                wanted.append((op, value))
                break

    if not wanted:
        return [], True
    column = target or _money_numeric_column(headers, rows, numeric_cols) or (
        numeric_cols[0] if len(numeric_cols) == 1 else None)
    if column is None:
        return [], False
    for op, value in wanted:
        filters.append({"column": column, "op": op, "value": value})
    return filters, True


def _whole_cell_filter(question: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
                       numeric_cols: Sequence[str], exclude: Sequence[str]) -> dict | None:
    """A filter on a cell the question quotes in full ("the debit of 03 Sep").

    Word-by-word matching cannot pick a row whose label is a date or a code, because each
    part on its own ("03", "sep") matches every row. A cell whose entire text appears in the
    question identifies exactly one row, so that is the strongest filter there is.
    """
    asked = _norm_text(question).strip()
    best: tuple[int, dict] | None = None
    for index, name in enumerate(headers):
        if name in exclude or name in numeric_cols:
            continue
        for row in rows:
            cell = _norm_text(_cell(row, index)).strip()
            # Only labels carrying a number ("03 sep", "inv-1003", "q2 2026") are matched whole:
            # a plain name is already handled, one word at a time, by the text filters.
            if len(cell) < 3 or cell in _STOP or not any(ch.isdigit() for ch in cell):
                continue
            if re.search(rf"(?<![a-z0-9]){re.escape(cell)}(?![a-z0-9])", asked):
                candidate = {"column": name, "op": "equals", "value": _squash(_cell(row, index))}
                if best is None or len(cell) > best[0]:
                    best = (len(cell), candidate)
    return best[1] if best else None


def _text_filters(question: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
                  numeric_cols: Sequence[str], exclude: Sequence[str],
                  skip_words: Iterable[str] = ()) -> list[dict]:
    """Text filters found by matching the table's own cell values against the question.

    Rather than guessing from phrasing ("from X" could be a vendor, a city or a
    date), this looks at what is actually in each non-numeric column: if a word of
    some cell value appears in the question, that column gets a ``contains`` filter
    on that word. "invoices from ABC" over a *Vendor* column holding "ABC Traders"
    yields ``Vendor contains 'abc'`` - no phrase parsing, and impossible to invent a
    value the table does not have.

    The most selective match per column wins, at most two columns are filtered, and
    intent words are never candidates.
    """
    qwords = set(_words(question))
    skip = {w.lower() for w in skip_words}
    candidates: list[tuple[float, dict]] = []
    for index, name in enumerate(headers):
        if name in exclude or name in numeric_cols:
            continue
        header_words = set(_words(name))
        best: tuple[float, str] | None = None
        for row in rows:
            cell = _cell(row, index)
            cell_exact = _norm_text(cell)
            for word in _words(cell):
                # "tables" is normally an intent word, but when a cell IS "Tables" the question
                # is naming that row, so an exact whole-cell match is allowed to filter.
                intent_only = word in _STOP and cell_exact != _norm_text(word)
                if len(word) < 3 or intent_only or word in skip or word in header_words:
                    continue
                if word not in qwords or parse_number(word) is not None:
                    continue
                hits = sum(1 for r in rows if re.search(rf"\b{re.escape(word)}\b", _norm_text(_cell(r, index))))
                if not hits:
                    continue
                selectivity = 1.0 - hits / max(1, len(rows))
                score = len(word) + 4.0 * selectivity
                if best is None or score > best[0]:
                    best = (score, word)
        if best:
            candidates.append((best[0], {"column": name, "op": "contains", "value": best[1]}))
    candidates.sort(key=lambda pair: (-pair[0], pair[1]["column"]))
    return [f for _, f in candidates[:2]]


def _row_matches(row: Sequence[str], headers: Sequence[str], f: dict) -> bool:
    """Apply one filter dict to one row."""
    index = _col_index(headers, f["column"])
    if index < 0:
        return False
    cell = _cell(row, index)
    op = f["op"]
    if op == "equals":
        return _norm_text(cell) == _norm_text(str(f["value"]))
    if op == "contains":
        return bool(re.search(rf"\b{re.escape(str(f['value']).lower())}\b", _norm_text(cell)))
    if op in ("month", "year", "quarter", "fy"):
        return _match_period(cell, op, str(f["value"]))
    number = parse_number(cell)
    if number is None:
        return False
    if op == "between":
        low, high = f["value"]
        return low <= number <= high
    target = float(f["value"])
    if op == ">":
        return number > target
    if op == ">=":
        return number >= target
    if op == "<":
        return number < target
    if op == "<=":
        return number <= target
    if op in ("==", "="):
        return number == target
    return False


_PERIODISH = re.compile(r"^(?:q[1-4]|fy\d{0,4}|\d+)$")


def _unresolved_entity(question: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
                       filters: Sequence[dict]) -> str | None:
    """A proper noun in the question that this table cannot account for, if any.

    "the total amount from LMN Corp" over a register that has no LMN row would
    otherwise be answered with the grand total of every vendor: arithmetic that is
    correct for a question nobody asked. So any capitalised word that is not a
    column name, not a value somewhere in the table and not part of a period is
    taken as proof that the question is about something this table does not hold,
    and the answer is refused.

    Only capitalised words are checked - a lower-cased question carries no signal
    about which of its words are names - and the first word is skipped because its
    capital is just a sentence start.
    """
    covered: set[str] = set()
    for f in filters:
        if isinstance(f.get("value"), str):
            covered.update(_words(f["value"]))
    for header in headers:
        covered.update(_words(header))
    for row in rows[:2000]:
        for cell in row:
            covered.update(_words(cell))
    for position, token in enumerate(str(question).split()):
        if position == 0 or not token[:1].isupper():
            continue
        for word in _words(token):
            if len(word) < 3 or word in _STOP or word in covered:
                continue
            if word in _MONTH_NAMES or _PERIODISH.match(word):
                continue
            return word
    return None


def _describe_filter(f: dict) -> str:
    """One filter as the clause that goes into `workings`."""
    op = f["op"]
    column, value = f["column"], f["value"]
    if op == "equals":
        return f"{column} is '{value}'"
    if op == "contains":
        return f"{column} contains '{value}'"
    if op == "month":
        return f"{column} = '{value}'"
    if op == "year":
        return f"{column} year = {value}"
    if op in ("quarter", "fy"):
        return f"{column} in {value}"
    if op == "between":
        return f"{column} between {_plain(value[0])} and {_plain(value[1])}"
    return f"{column} {op} {_plain(float(value))}"


# ================================================================ formatting
def _indian_group(digits: str) -> str:
    """"235000" -> "2,35,000": last three digits, then pairs."""
    if len(digits) <= 3:
        return digits
    head, tail = digits[:-3], digits[-3:]
    parts: list[str] = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts + [tail])


def format_indian(value: float, decimals: int = 2) -> str:
    """Indian digit grouping: 235000 -> "2,35,000.00", 1500000 -> "15,00,000.00"."""
    sign = "-" if value < 0 else ""
    text = f"{abs(float(value)):.{decimals}f}"
    whole, _, frac = text.partition(".")
    grouped = _indian_group(whole)
    return f"{sign}{grouped}.{frac}" if frac else f"{sign}{grouped}"


def _plain(value: float) -> str:
    """A plain number with up to two decimals and no trailing zero noise."""
    if value == int(value) and abs(value) < 1e15:
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def _format_value(value: float, operation: str, money: bool, row_count: int,
                  grouped_source: bool = False) -> str:
    """A total reads the way the column reads.

    ``money`` means the question or the header said so, and gets the rupee sign.
    ``grouped_source`` means the cells themselves are written with thousands separators
    (57,530 · 6,20,000), so the total is grouped the same way — "421070" is the same number
    as "4,21,070", but only one of them is readable.
    """
    if operation == "count":
        return f"{row_count} row" if row_count == 1 else f"{row_count} rows"
    if money:
        return "₹" + format_indian(value)
    if grouped_source and abs(value) >= 10000:
        return format_indian(value)
    return _plain(value)


# ========================================================= query answering
def _score_table(t: dict, question: str) -> float:
    """How well a table matches the question, on words alone - no embeddings here.

    Header words count most (they are what a filter or an aggregate will name), the
    title next, and a cell value that appears in the question last but still enough
    to pull the right table out of a pile ("ABC Traders" -> the purchases table).
    """
    qwords = set(w for w in _words(question) if w not in _STOP)
    if not qwords:
        return 0.0
    score = 0.0
    for header in t.get("headers") or []:
        score += 3.0 * len(qwords & {w for w in _words(header) if w not in _STOP})
    score += 2.0 * len(qwords & {w for w in _words(t.get("title") or "") if w not in _STOP})
    cell_words: set[str] = set()
    for row in (t.get("rows") or [])[:200]:
        for cell in row:
            cell_words.update(w for w in _words(cell) if len(w) >= 3 and w not in _STOP)
    score += 1.0 * len(qwords & cell_words)
    return score


def _numeric_cols_of(t: dict) -> list[str]:
    """The table's numeric columns, always recomputed from the rows.

    The stored list is only a cache written when the table was indexed, and an older build
    wrote the wrong one: a Credit column full of em-dashes was not recognised as numeric, so
    "total credit" answered "no column could be identified" on a table that plainly has one.
    Detection over a handful of rows costs nothing, so the live answer never depends on what
    an old index happened to record. The union keeps any column the stored list knew about.
    """
    headers = [str(h) for h in t.get("headers") or []]
    rows = [list(r) for r in t.get("rows") or []]
    live = detect_numeric_cols(headers, rows)
    stored = [c for c in (t.get("numeric_cols") or []) if c in headers and c not in live]
    values = lambda name: [_squash(r[headers.index(name)] if headers.index(name) < len(r) else "") for r in rows]
    return live + [c for c in stored if not _is_index_column(c, values(c))]


_TOTAL_LABELS = frozenset({"total", "totals", "grand total", "grand-total", "sub total", "subtotal",
                           "sub-total", "net total", "overall total", "sum", "total:", "all"})


def _is_total_row(row: Sequence[str], headers: Sequence[str]) -> bool:
    """Is this the table's own summary row?

    A table that ends in a TOTAL row already contains its own sum. Adding that row into a
    sum counts every amount twice - the reason "what is the total" once answered 1,44,900
    for a table whose printed total was 72,450.
    """
    for cell in row:
        text = _squash(cell)
        if not text:
            continue
        return _norm_text(text).strip(" :") in _TOTAL_LABELS
    return False


def _answer_from_table(question: str, t: dict, operation: str) -> dict | None:
    """Try to answer from one table. ``None`` means "this table cannot answer it"."""
    headers = [str(h) for h in t.get("headers") or []]
    rows = [list(r) for r in t.get("rows") or []]
    if not headers or not rows:
        return None
    numeric_cols = _numeric_cols_of(t)

    # The column the question names, if any. `count` needs one only to attach a
    # numeric comparison to ("how many invoices with quantity between 5 and 10") -
    # the count itself is always of rows.
    # "What is the count of Documents?" over a table with a *column* called Count is not a request
    # to count rows: the word names the column, and the answer is that column's value for the row
    # the question identifies. The same trap exists for a column called Total or Number.
    if operation == "count":
        named = [c for c in numeric_cols if any(w in _words(question) for w in _words(c))]
        if named:
            operation = "lookup"

    column, column_score = _resolve_numeric_column(question, headers, rows, numeric_cols)
    if column is None and operation not in ("count", "filter"):
        return None                         # no column to compute on: refuse

    filters: list[dict] = []
    numeric, ok = _numeric_filters(question, headers, rows, numeric_cols, column)
    if not ok:
        return None                         # "above 50000" with no column it can mean
    filters += numeric
    if numeric and column is None:
        column = numeric[0]["column"]       # the comparison told us which column is meant

    period = _period_filter(question, headers, rows, exclude=[column] if column else [])
    exact = _whole_cell_filter(question, headers, rows, numeric_cols, exclude=[column] if column else [])
    if exact:
        # The question quoted a whole labelled cell ("the debit of 03 Sep"), which picks one row.
        # That is more specific than "September", so it wins over a period filter on that column.
        filters.append(exact)
        if period and period["column"] != exact["column"]:
            filters.append(period)
    elif period:
        filters.append(period)
    skip_words = set(_words(period["value"])) if period else set()
    # "total credit" names the Credit COLUMN; it must not also filter rows whose text
    # happens to contain "credit" ("SALARY CREDIT"), which would silently drop rows.
    for header in headers:
        if column and _norm_text(header) == _norm_text(column):
            skip_words |= {w for w in _words(header) if len(w) >= 3}
    filters += _text_filters(
        question, headers, rows, numeric_cols,
        exclude=[f["column"] for f in filters] + ([column] if column else []),
        skip_words=skip_words,
    )

    if operation == "filter" and not filters:
        return None                         # a lookup with nothing to look up
    if _unresolved_entity(question, headers, rows, filters):
        return None                         # the question names something not in this table

    matched = [r for r in rows if all(_row_matches(r, headers, f) for f in filters)]
    # "How many scanned pages?" against a table of categories and counts is asking for the
    # number in that row, not for how many rows carry the label.
    # "What is the height of basil?" wants the value in that column; "list the invoices" and
    # "how many invoices" are asking for rows, and stay as they are.
    asks_for_rows = re.search(r"\b(list|show|which|how many|count of rows|rows)\b", question, re.I)
    # "How many scanned pages?" against a table whose one numeric column IS a count, with a row
    # named by the question, is asking for that number - not for how many rows carry the label.
    label_filter = any(f["op"] in ("contains", "equals") for f in filters)
    category_count = len(numeric_cols) == 1 and label_filter and len(matched) == 1
    # "How many labour?" against a parts table names a row, not a set of rows to count.
    # Answering "1 row" is useless - what was asked for is that row's quantity. Only a
    # quantity column can answer it; a table without one falls through to the row count.
    if operation == "count" and column is None and label_filter and len(matched) == 1:
        quantity_col = next((h for h in headers
                             if _norm_text(h) in _QTY_HEADER_WORDS
                             or any(w in _QTY_HEADER_WORDS for w in _words(h))), None)
        if quantity_col:
            column, operation = quantity_col, "lookup"
    if (operation in ("count", "filter") and column and matched and len(matched) < len(rows)
            and (column_score >= 4 or len(matched) == 1)
            and (not asks_for_rows or category_count)):
        index = _col_index(headers, column)
        if all(parse_number(_cell(r, index)) is not None for r in matched):
            operation = "lookup"
    if not matched:
        # An empty result is more often a mis-read filter than a true zero, so the
        # caller falls back to retrieval instead of asserting "nothing".
        return None

    # Drop the table's own TOTAL row before doing arithmetic over the column, unless the
    # question asked for that row itself ("what does the total row say").
    dropped_total_row = False
    if operation in ("sum", "avg", "max", "min", "diff", "count") and len(matched) > 1:
        data_rows = [r for r in matched if not _is_total_row(r, headers)]
        if data_rows and len(data_rows) < len(matched):
            matched, dropped_total_row = data_rows, True

    value: float | None = None
    evidence = matched
    if operation == "lookup":
        # The value of one column for the rows the question identified. With several rows
        # they are added, which is what "the count of documents and images" means.
        if column is None:
            return None
        index = _col_index(headers, column)
        usable = [(parse_number(_cell(r, index)), r) for r in matched]
        usable = [(n, r) for n, r in usable if n is not None]
        blank_cell = False
        if not usable:
            cells = [_squash(_cell(r, index)) for r in matched]
            # The row exists and the cell is empty or a dash: that is a real answer - none -
            # and reporting 0 is more use than falling back to a paragraph of prose.
            if cells and all(c == "" or c.lower() in _EMPTY_CELLS for c in cells):
                blank_cell = True
                evidence, row_count, result_values, value = matched, len(matched), None, 0.0
            else:
                return None
        if not blank_cell:
            evidence = [r for _, r in usable]
            row_count = len(usable)
            # Two counts of the same kind add up; two measurements do not. "The height of basil"
            # over a log with two basil rows is 18 cm and 21 cm, not 39 cm.
            cells = [_squash(_cell(r, index)) for r in evidence]
            measured = any(re.search(r"\d\s*[A-Za-z°%]", c) and not _MONEY_MARKER.search(c) for c in cells)
            if measured:
                result_values = [c for c in cells if c]
                value = float(usable[0][0])
            else:
                value = float(sum(n for n, _ in usable))
                result_values = None
    elif operation in ("sum", "avg", "max", "min", "diff"):
        assert column is not None
        index = _col_index(headers, column)
        pairs = [(parse_number(_cell(r, index)), r) for r in matched]
        usable = [(n, r) for n, r in pairs if n is not None]
        if not usable:
            return None
        numbers = [n for n, _ in usable]
        evidence = [r for _, r in usable]
        if operation == "sum":
            value = float(sum(numbers))
        elif operation == "avg":
            value = float(sum(numbers)) / len(numbers)
        elif operation == "max":
            value = max(numbers)
            evidence = [max(usable, key=lambda pair: pair[0])[1]]
        elif operation == "min":
            value = min(numbers)
            evidence = [min(usable, key=lambda pair: pair[0])[1]]
        else:
            value = max(numbers) - min(numbers)
        row_count = len(usable)
    elif operation == "count":
        row_count = len(matched)
        value = float(row_count)
    else:                                   # filter / lookup
        row_count = len(matched)

    where = " and ".join(_describe_filter(f) for f in filters)
    without_total = " (excluding the table's own TOTAL row)" if locals().get("dropped_total_row") else ""
    if operation == "lookup" and locals().get("blank_cell"):
        workings = f"{column}{' where ' + where if where else ''} is blank in the table, so 0"
    elif operation == "lookup":
        workings = f"{column}{' where ' + where if where else ''} = {value}"
    elif operation == "count":
        workings = f"count(rows){' where ' + where if where else ''} = {row_count}{without_total}"
    elif operation == "filter":
        workings = f"rows where {where} ({row_count} matched)"
    elif operation == "diff":
        workings = (f"max({column}) - min({column}) over {row_count} rows"
                    f"{' where ' + where if where else ''}")
    else:
        name = {"sum": "sum", "avg": "avg", "max": "max", "min": "min"}.get(operation, operation)
        workings = (f"{name}({column}) over {row_count} rows"
                    f"{' where ' + where if where else ''}{without_total}")

    money = bool(column) and _column_is_money(headers, rows, column, question)
    # "high" needs the question to have actually named the column (by header word or
    # by a money/quantity synonym), or - for a lookup - a filter matched against a
    # real cell value. The bare "one numeric column, so it must be that one"
    # fallback scores 1 and stays "medium".
    named_column = column_score >= 3
    matched_by_value = any(f["op"] == "contains" for f in filters)
    # A value the question did not ask for is not an answer: "Who is Ananya Rao?" must not be
    # answered with a number just because her name appears in a cell. Only a plain row lookup
    # may rest on a matched value alone.
    single_column = len(numeric_cols) == 1
    confident = (named_column
                 or (operation == "filter" and matched_by_value)
                 # one numeric column and a row the question named: "how many scanned pages"
                 # over a table of categories and counts has only one possible reading
                 or (operation == "lookup" and matched_by_value and single_column))
    result: dict = {
        "table_id": t.get("id"),
        "doc_id": t.get("doc_id"),
        "doc_name": t.get("doc_name"),
        "page": t.get("page"),
        "title": t.get("title", ""),
        "operation": operation,
        "column": column,
        "filters": filters,
        "matched_rows": [dict(zip(headers, list(r) + [""] * (len(headers) - len(r))))
                         for r in evidence[:MAX_MATCHED_ROWS]],
        "row_count": row_count,
        "workings": workings,
        "confidence": "high" if confident else "medium",
    }
    if operation == "lookup" and locals().get("result_values"):
        result["values"] = result_values
    if operation != "filter":
        result["value"] = float(value if value is not None else 0.0)
        grouped = False
        if column and column in headers:
            index = headers.index(column)
            grouped = any("," in _cell(r, index) for r in rows)
        if result.get("values"):
            listed = result["values"]
            result["formatted"] = " and ".join(listed[:4]) + (" …" if len(listed) > 4 else "")
        else:
            result["formatted"] = _format_value(result["value"], operation, money, row_count, grouped)
            # "31" answers "the tallest plant" less well than "31 cm": carry the column's unit.
            if column and operation in ("max", "min", "avg", "sum", "diff"):
                index = _col_index(headers, column)
                units = {m.group(1) for m in (re.search(r"\d\s*([A-Za-z°%µ/]{1,6})$", _squash(_cell(r, index)))
                                              for r in rows) if m}
                if len(units) == 1:
                    unit = units.pop()
                    if unit.lower() not in _MONTH_NAMES:
                        result["formatted"] = f"{result['formatted']} {unit}"
    else:
        result["formatted"] = f"{row_count} row" if row_count == 1 else f"{row_count} rows"
    return result


def answer_table_query(question: str, tables: list[dict], strict: bool = False) -> dict | None:
    """Answer a numeric table question with Python arithmetic, or return ``None``.

    The flow is: work out the operation from the wording, rank the tables by plain
    word overlap with the question, then for each of the best few tables try to pin
    down a numeric column and the filters that narrow the rows, and compute.

    ``None`` is returned - and it is returned often, on purpose - whenever the
    question is not an arithmetic question, no numeric column can be identified
    without guessing, a comparison cannot be attached to a column, or no row matches
    the filters. In all of those cases the caller should fall back to ordinary
    retrieval: an honest miss is cheap, a confidently wrong total is not.

    The returned dict carries the source keys of the table it used (``table_id``,
    ``doc_id``, ``doc_name``, ``page``, ``title``), the operation and column, the
    filters it applied, the rows it used as evidence (at most 50), the computed
    ``value``, a display string in ``formatted`` and a one-line ``workings``
    explanation that can be shown to the user verbatim. A pure lookup has operation
    ``"filter"`` and no ``value``.
    """
    if not tables:
        return None
    question = str(question or "").strip()
    if not question:
        return None
    operation = _parse_operation(question)
    if operation is None:
        # Inside one table a bare phrase is a question: "height of basil" means "what is the
        # height of basil". Across the whole workspace that is too loose, so it still needs a verb.
        if strict:
            return None
        operation = "filter"

    scored = sorted(((_score_table(t, question), i, t) for i, t in enumerate(tables)),
                    key=lambda triple: (-triple[0], triple[1]))
    candidates = [(score, t) for score, _, t in scored if score > 0][:3]
    if not candidates and len(tables) == 1:
        candidates = [(0.0, tables[0])]     # the caller already narrowed it to one table
    for _score, t in candidates:
        answer = _answer_from_table(question, t, operation)
        if answer is None:
            continue
        # Asked across every table in the workspace, a guess is worse than a miss: the question
        # has to have named the column or matched a real cell. Inside one table the user has
        # already said which table they mean, so a weaker match is fine there.
        if strict and answer.get("confidence") != "high":
            continue
        if strict and answer.get("operation") == "filter":
            # "1 matching row" is not an answer to a question asked of the whole workspace;
            # ordinary retrieval will say what that row actually contains.
            continue
        return answer
    return None
